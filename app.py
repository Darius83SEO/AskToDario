"""
AskToDario - People Also Ask Explorer
by Dario Ciraci, Senior SEO Consultant (webinfermento.it)
"""

import streamlit as st
import graphviz
import pandas as pd
import time
import csv
import io
import base64
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sentence_transformers import SentenceTransformer, util
import torch

from scraper import (
    get_paa_cached, extract_paa_tree,
    LANGUAGES, COUNTRIES, BRANCH_COLORS
)

# === PAGE CONFIG ===
st.set_page_config(
    page_title="AskToDario - PAA Explorer",
    page_icon="💬",
    layout="wide",
    initial_sidebar_state="expanded"
)

# === LOAD AVATAR ===
def get_avatar_base64():
    avatar_path = os.path.join(os.path.dirname(__file__), "assets", "avatar.png")
    if os.path.exists(avatar_path):
        with open(avatar_path, "rb") as f:
            return base64.b64encode(f.read()).decode()
    return ""

avatar_b64 = get_avatar_base64()

# === CSS ===
st.markdown("""
<style>
    /* General */
    .block-container { padding-top: 1rem; max-width: 1200px; }
    h1, h2, h3 { color: #1A1A2E; }

    /* Logo header */
    .logo-container {
        display: flex;
        align-items: center;
        justify-content: center;
        gap: 20px;
        padding: 25px 0 10px 0;
    }
    .logo-avatar {
        width: 80px;
        height: 80px;
        border-radius: 50%;
        border: 3px solid #4A6CF7;
        box-shadow: 0 4px 15px rgba(74, 108, 247, 0.3);
    }
    .logo-text {
        font-size: 2.8rem;
        font-weight: 800;
        background: linear-gradient(135deg, #4A6CF7 0%, #6C63FF 50%, #4A6CF7 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        letter-spacing: -1px;
    }
    .logo-subtitle {
        text-align: center;
        color: #7A7A9E;
        font-size: 1.05rem;
        margin-top: -8px;
        margin-bottom: 25px;
    }

    /* Cards */
    .stat-card {
        background: linear-gradient(135deg, #F4F6FC 0%, #FFFFFF 100%);
        border: 1px solid #E8ECF4;
        border-radius: 12px;
        padding: 18px;
        text-align: center;
    }
    .stat-number {
        font-size: 2.2rem;
        font-weight: 700;
        color: #4A6CF7;
    }
    .stat-label {
        color: #7A7A9E;
        font-size: 0.85rem;
        text-transform: uppercase;
        letter-spacing: 1px;
    }

    /* PAA cards */
    .paa-branch {
        border-left: 4px solid;
        padding: 12px 16px;
        margin: 6px 0;
        border-radius: 0 10px 10px 0;
        background: #F8F9FE;
        transition: background 0.2s;
    }
    .paa-branch:hover { background: #EEF0FA; }
    .paa-branch-title {
        font-weight: 700;
        font-size: 1.05rem;
    }
    .paa-sub {
        color: #4A4A6A;
        padding: 3px 0 3px 20px;
        font-size: 0.95rem;
    }

    /* Cluster */
    .cluster-badge {
        display: inline-block;
        background: #4A6CF7;
        color: white;
        padding: 3px 12px;
        border-radius: 20px;
        font-size: 0.8rem;
        font-weight: 600;
    }

    /* Sidebar */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #F4F6FC 0%, #FFFFFF 100%);
    }
    [data-testid="stSidebar"] h2 {
        color: #4A6CF7;
    }

    /* Search box */
    .search-container {
        background: linear-gradient(135deg, #4A6CF7 0%, #6C63FF 100%);
        padding: 30px;
        border-radius: 16px;
        margin-bottom: 25px;
    }
    .search-container h3 {
        color: white !important;
        margin-bottom: 15px;
    }

    /* Footer */
    .footer {
        text-align: center;
        padding: 30px 0 10px 0;
        color: #B0B0C8;
        font-size: 0.85rem;
        border-top: 1px solid #E8ECF4;
        margin-top: 40px;
    }
    .footer a { color: #4A6CF7 !important; text-decoration: none; }
</style>
""", unsafe_allow_html=True)

# === NLP MODEL ===
@st.cache_resource
def load_nlp_model():
    return SentenceTransformer('paraphrase-MiniLM-L6-v2')

nlp_model = load_nlp_model()


# === LOGO HEADER ===
if avatar_b64:
    st.markdown(f"""
    <div class="logo-container">
        <img src="data:image/png;base64,{avatar_b64}" class="logo-avatar" />
        <span class="logo-text">AskToDario</span>
    </div>
    """, unsafe_allow_html=True)
else:
    st.markdown('<div class="logo-container"><span class="logo-text">💬 AskToDario</span></div>',
                unsafe_allow_html=True)

st.markdown('<p class="logo-subtitle">People Also Ask Explorer — Scopri cosa chiedono gli utenti su Google</p>',
            unsafe_allow_html=True)


# === SIDEBAR ===
with st.sidebar:
    if avatar_b64:
        st.markdown(f"""
        <div style="text-align:center; padding: 10px 0 20px 0;">
            <img src="data:image/png;base64,{avatar_b64}"
                 style="width:60px; height:60px; border-radius:50%; border:2px solid #4A6CF7;" />
            <div style="font-weight:700; color:#4A6CF7; margin-top:5px;">AskToDario</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("## Impostazioni")

    lang_name = st.selectbox("🌐 Lingua ricerca", list(LANGUAGES.keys()), index=0)
    country_name = st.selectbox("📍 Paese Google", list(COUNTRIES.keys()), index=0)
    depth = st.slider("🔽 Profondita' estrazione", min_value=1, max_value=3, value=2,
                       help="1 = 4 PAA, 2 = ~20 PAA, 3 = ~84 PAA (piu' lento)")

    hl = LANGUAGES[lang_name]
    gl, google_domain = COUNTRIES[country_name]

    est_requests = sum(4**i for i in range(1, depth + 1)) + 1
    est_time = int(est_requests * 1.5)
    st.caption(f"Richieste stimate: ~{est_requests} | Tempo: ~{est_time//60}m {est_time%60}s")

    st.markdown("---")
    st.markdown(
        "**Creato da [Dario Ciraci](https://www.webinfermento.it/consulente-seo/)**\n\n"
        "Senior SEO Consultant"
    )


# === TABS ===
tab_search, tab_compare = st.tabs(["🔍 Ricerca PAA", "📊 Confronta nel tempo"])


# ==========================================
# TAB 1: RICERCA
# ==========================================
with tab_search:
    # Search area
    st.markdown("""
    <div class="search-container">
        <h3>🔍 Cerca People Also Ask</h3>
    </div>
    """, unsafe_allow_html=True)

    col1, col2 = st.columns([5, 1])
    with col1:
        query = st.text_input("Keyword o domanda:",
                              placeholder="es. pizza napoletana, come fare il pane, SEO strategy...",
                              label_visibility="collapsed")
    with col2:
        search_btn = st.button("Cerca", type="primary", use_container_width=True)

    # === ESECUZIONE RICERCA ===
    if search_btn and query.strip():
        query = query.strip()

        progress_bar = st.progress(0, text="Avvio estrazione PAA...")
        status_text = st.empty()

        def update_progress(current, total, question):
            pct = min(current / max(total, 1), 1.0)
            progress_bar.progress(pct, text=f"[{current}/{total}] {question[:60]}...")

        with st.spinner(""):
            all_questions, edges, question_branches = extract_paa_tree(
                query, hl=hl, gl=gl, google_domain=google_domain,
                depth=depth, progress_callback=update_progress
            )

        progress_bar.progress(1.0, text="Completato!")
        time.sleep(0.3)
        progress_bar.empty()
        status_text.empty()

        st.session_state['results'] = {
            'query': query, 'all_questions': all_questions,
            'edges': edges, 'question_branches': question_branches,
            'hl': hl, 'gl': gl, 'google_domain': google_domain,
        }

    # === MOSTRA RISULTATI ===
    if 'results' in st.session_state:
        data = st.session_state['results']
        all_questions = data['all_questions']
        edges = data['edges']
        question_branches = data['question_branches']
        query = data['query']

        paa_only = [q for q in all_questions if q[1] > 0]

        if not paa_only:
            st.warning("Nessun PAA trovato per questa query. Prova con un'altra keyword o controlla lingua/paese.")
        else:
            # --- Stats ---
            st.markdown("### 📈 Risultati")
            c1, c2, c3, c4 = st.columns(4)
            with c1:
                st.markdown(f'<div class="stat-card"><div class="stat-number">{len(paa_only)}</div>'
                           f'<div class="stat-label">PAA totali</div></div>', unsafe_allow_html=True)
            with c2:
                lvl1 = sum(1 for q in all_questions if q[1] == 1)
                st.markdown(f'<div class="stat-card"><div class="stat-number">{lvl1}</div>'
                           f'<div class="stat-label">Livello 1</div></div>', unsafe_allow_html=True)
            with c3:
                lvl2 = sum(1 for q in all_questions if q[1] == 2)
                st.markdown(f'<div class="stat-card"><div class="stat-number">{lvl2}</div>'
                           f'<div class="stat-label">Livello 2</div></div>', unsafe_allow_html=True)
            with c4:
                lvl3 = sum(1 for q in all_questions if q[1] == 3)
                st.markdown(f'<div class="stat-card"><div class="stat-number">{lvl3}</div>'
                           f'<div class="stat-label">Livello 3</div></div>', unsafe_allow_html=True)

            # --- ALBERO ---
            st.markdown("### 🌳 Albero PAA")
            st.caption("Clicca e trascina per navigare. La keyword parte da sinistra, i PAA si espandono verso destra.")

            dot = graphviz.Digraph(
                graph_attr={
                    'rankdir': 'LR',
                    'splines': 'ortho',
                    'nodesep': '0.25',
                    'ranksep': '0.9',
                    'bgcolor': 'transparent',
                    'dpi': '120',
                    'pad': '0.3',
                },
                node_attr={
                    'shape': 'box',
                    'style': 'rounded,filled',
                    'fontname': 'Arial',
                    'fontsize': '10',
                    'margin': '0.12,0.06',
                    'color': '#E0E0E0',
                },
                edge_attr={
                    'arrowsize': '0.5',
                    'penwidth': '1.8',
                }
            )

            def wrap_text(text, max_len=38):
                if len(text) <= max_len:
                    return text
                words = text.split()
                lines, current = [], ""
                for w in words:
                    if len(current) + len(w) + 1 > max_len:
                        lines.append(current)
                        current = w
                    else:
                        current = (current + " " + w).strip()
                if current:
                    lines.append(current)
                return "\\n".join(lines[:3])

            for text, level, parent, branch_idx in all_questions:
                nid = str(hash(text))
                label = wrap_text(text)
                if level == 0:
                    dot.node(nid, label, fillcolor='#4A6CF7', fontcolor='white',
                             fontsize='13', penwidth='2', color='#4A6CF7')
                else:
                    color = BRANCH_COLORS[branch_idx % len(BRANCH_COLORS)]
                    if level == 1:
                        dot.node(nid, label, fillcolor=color, fontcolor='white',
                                 fontsize='11', penwidth='2', color=color)
                    elif level == 2:
                        dot.node(nid, label, fillcolor=color + '30',
                                 fontcolor='#333333', fontsize='10',
                                 penwidth='1.2', color=color + '80')
                    else:
                        dot.node(nid, label, fillcolor=color + '15',
                                 fontcolor='#555555', fontsize='9',
                                 penwidth='1', color=color + '50')

            for parent_t, child_t in edges:
                branch = question_branches.get(child_t, 0)
                color = BRANCH_COLORS[branch % len(BRANCH_COLORS)]
                dot.edge(str(hash(parent_t)), str(hash(child_t)), color=color + 'AA')

            st.graphviz_chart(dot, use_container_width=True)

            # --- LISTA PAA per ramo ---
            st.markdown("### 📋 PAA per ramo")

            for branch_idx in range(min(4, max((q[3] for q in all_questions if q[1] > 0), default=0) + 1)):
                branch_qs = [q for q in all_questions if q[3] == branch_idx]
                if not branch_qs:
                    continue
                color = BRANCH_COLORS[branch_idx % len(BRANCH_COLORS)]
                root_q = next((q for q in branch_qs if q[1] == 1), None)
                if root_q:
                    st.markdown(
                        f'<div class="paa-branch" style="border-color: {color};">'
                        f'<span class="paa-branch-title" style="color: {color};">'
                        f'Ramo {branch_idx+1}: {root_q[0]}</span></div>',
                        unsafe_allow_html=True
                    )
                    sub_qs = [q for q in branch_qs if q[1] > 1]
                    for q in sub_qs:
                        prefix = "└" if q == sub_qs[-1] else "├"
                        indent = "&nbsp;" * ((q[1] - 1) * 6)
                        st.markdown(
                            f'<div class="paa-sub">{indent}{prefix} {q[0]}</div>',
                            unsafe_allow_html=True
                        )

            # --- CLUSTERING ---
            if len(paa_only) >= 4:
                st.markdown("### 🧠 Clustering Semantico")
                st.caption("Domande raggruppate per similarita' di intento tramite NLP.")

                with st.spinner("Analisi semantica in corso..."):
                    texts = [q[0] for q in paa_only]
                    embeddings = nlp_model.encode(texts, convert_to_tensor=True, show_progress_bar=False)
                    sim_matrix = util.pytorch_cos_sim(embeddings, embeddings).cpu().numpy()

                    clusters = {}
                    assigned = set()
                    for i in range(len(texts)):
                        if i in assigned:
                            continue
                        cluster = [texts[i]]
                        assigned.add(i)
                        for j in range(i + 1, len(texts)):
                            if j not in assigned and sim_matrix[i][j] > 0.55:
                                cluster.append(texts[j])
                                assigned.add(j)
                        if len(cluster) >= 2:
                            name = min(cluster, key=len)
                            clusters[name] = cluster

                    unclustered = [texts[i] for i in range(len(texts)) if i not in assigned]
                    if unclustered:
                        clusters["Altre domande"] = unclustered

                cols_cl = st.columns(min(len(clusters), 3))
                for idx, (name, qs) in enumerate(clusters.items()):
                    with cols_cl[idx % len(cols_cl)]:
                        st.markdown(f'<span class="cluster-badge">{len(qs)} domande</span>',
                                   unsafe_allow_html=True)
                        with st.expander(f"**{name}**", expanded=idx < 3):
                            for q in qs:
                                st.write(f"- {q}")

            # --- EXPORT ---
            st.markdown("### 💾 Esporta risultati")
            col_e1, col_e2, col_e3 = st.columns(3)

            with col_e1:
                wb = Workbook()
                ws = wb.active
                ws.title = "PAA Tree"
                ws.append(["AskToDario - PAA Extractor by Dario Ciraci"])
                ws.cell(row=1, column=1).font = Font(size=14, bold=True)
                ws.append(["webinfermento.it/consulente-seo/"])
                ws.append([""])
                ws.append(["Domanda", "Livello", "Domanda padre", "Ramo"])
                for col in range(1, 5):
                    ws.cell(row=4, column=col).font = Font(bold=True, color="FFFFFF")
                    ws.cell(row=4, column=col).fill = PatternFill(
                        start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
                blue_fill = PatternFill(start_color="DAE2FC", end_color="DAE2FC", fill_type="solid")
                for text, level, parent, branch in all_questions:
                    ws.append([text, level, parent or "",
                              f"Ramo {branch+1}" if branch >= 0 else "Root"])
                    if level == 1:
                        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
                        ws.cell(row=ws.max_row, column=1).fill = blue_fill
                if clusters:
                    ws2 = wb.create_sheet("Cluster Semantici")
                    ws2.append(["Cluster", "Domanda"])
                    ws2.cell(row=1, column=1).font = Font(bold=True, color="FFFFFF")
                    ws2.cell(row=1, column=2).font = Font(bold=True, color="FFFFFF")
                    ws2.cell(row=1, column=1).fill = PatternFill(
                        start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
                    ws2.cell(row=1, column=2).fill = PatternFill(
                        start_color="4A6CF7", end_color="4A6CF7", fill_type="solid")
                    for cname, cqs in clusters.items():
                        for cq in cqs:
                            ws2.append([cname, cq])
                buf = io.BytesIO()
                wb.save(buf)
                st.download_button("📥 Scarica Excel", buf.getvalue(),
                                  file_name=f"paa_{query.replace(' ','_')}.xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  use_container_width=True)

            with col_e2:
                csv_buf = io.StringIO()
                writer = csv.writer(csv_buf)
                writer.writerow(["Domanda", "Livello", "Domanda padre", "Ramo"])
                for text, level, parent, branch in all_questions:
                    writer.writerow([text, level, parent or "",
                                    f"Ramo {branch+1}" if branch >= 0 else "Root"])
                st.download_button("📥 Scarica CSV", csv_buf.getvalue(),
                                  file_name=f"paa_{query.replace(' ','_')}.csv",
                                  mime="text/csv", use_container_width=True)

            with col_e3:
                all_text = "\n".join([q[0] for q in paa_only])
                st.download_button("📥 Copia lista TXT", all_text,
                                  file_name=f"paa_{query.replace(' ','_')}.txt",
                                  mime="text/plain", use_container_width=True)

            # --- Espandi singolo PAA ---
            st.markdown("### 🔄 Espandi un PAA")
            st.caption("Seleziona un PAA per usarlo come nuova keyword e estrarre i suoi sotto-PAA.")
            paa_options = ["-- Seleziona --"] + [q[0] for q in paa_only]
            selected_paa = st.selectbox("PAA da espandere:", paa_options, label_visibility="collapsed")

            if selected_paa != "-- Seleziona --":
                if st.button("🔄 Espandi questo PAA", type="primary"):
                    with st.spinner(f"Estrazione per: {selected_paa[:50]}..."):
                        sub_q, sub_e, sub_b = extract_paa_tree(
                            selected_paa,
                            hl=data.get('hl', 'it'),
                            gl=data.get('gl', 'it'),
                            google_domain=data.get('google_domain', 'google.it'),
                            depth=2
                        )
                    st.session_state['results'] = {
                        'query': selected_paa, 'all_questions': sub_q,
                        'edges': sub_e, 'question_branches': sub_b,
                        'hl': data.get('hl', 'it'), 'gl': data.get('gl', 'it'),
                        'google_domain': data.get('google_domain', 'google.it'),
                    }
                    st.rerun()


# ==========================================
# TAB 2: CONFRONTO
# ==========================================
with tab_compare:
    st.markdown("### 📊 Confronta PAA nel tempo")
    st.caption("Carica due file Excel scaricati in momenti diversi per scoprire nuovi PAA, "
               "PAA rimossi, nuovi intenti di ricerca e indice di volatilita'.")

    col_u1, col_u2 = st.columns(2)
    with col_u1:
        file_new = st.file_uploader("📄 File NUOVO (piu' recente)", type=['xlsx'], key='fnew')
    with col_u2:
        file_old = st.file_uploader("📄 File PRECEDENTE", type=['xlsx'], key='fold')

    if file_new and file_old:
        if st.button("🔄 Confronta", type="primary", use_container_width=True):
            df_new = pd.read_excel(file_new)
            df_old = pd.read_excel(file_old)

            new_set = set(str(x).strip() for x in df_new.iloc[3:, 0].dropna().tolist() if str(x).strip())
            old_set = set(str(x).strip() for x in df_old.iloc[3:, 0].dropna().tolist() if str(x).strip())

            added = list(new_set - old_set)
            removed = list(old_set - new_set)
            volatility = (len(added) + len(removed)) / max(len(old_set), 1) * 100

            c1, c2, c3 = st.columns(3)
            with c1:
                st.markdown(f'<div class="stat-card"><div class="stat-number" style="color:#34A853;">'
                           f'+{len(added)}</div><div class="stat-label">Nuovi PAA</div></div>',
                           unsafe_allow_html=True)
            with c2:
                st.markdown(f'<div class="stat-card"><div class="stat-number" style="color:#EA4335;">'
                           f'-{len(removed)}</div><div class="stat-label">PAA rimossi</div></div>',
                           unsafe_allow_html=True)
            with c3:
                vol_color = "#EA4335" if volatility > 30 else "#FBBC05" if volatility > 15 else "#34A853"
                st.markdown(f'<div class="stat-card"><div class="stat-number" style="color:{vol_color};">'
                           f'{volatility:.0f}%</div><div class="stat-label">Volatilita\'</div></div>',
                           unsafe_allow_html=True)

            # Nuovi intenti
            if added:
                old_list = list(old_set)
                if old_list:
                    with st.spinner("Analisi intenti con NLP..."):
                        added_emb = nlp_model.encode(added, convert_to_tensor=True)
                        old_emb = nlp_model.encode(old_list, convert_to_tensor=True)
                        new_intents = []
                        for i, paa in enumerate(added):
                            score = torch.max(util.pytorch_cos_sim(added_emb[i], old_emb)).item()
                            if score < 0.7:
                                new_intents.append(paa)

                    if new_intents:
                        st.markdown("#### 🆕 Nuovi intenti di ricerca")
                        st.caption("PAA semanticamente diversi da tutto il set precedente (similarita' < 70%)")
                        for ni in new_intents:
                            st.markdown(f"- 🔴 **{ni}**")

            col_a, col_r = st.columns(2)
            with col_a:
                st.markdown("#### ➕ PAA aggiunti")
                if added:
                    for a in added:
                        st.markdown(f"- {a}")
                else:
                    st.info("Nessun nuovo PAA")
            with col_r:
                st.markdown("#### ➖ PAA rimossi")
                if removed:
                    for r in removed:
                        st.markdown(f"- ~~{r}~~")
                else:
                    st.info("Nessun PAA rimosso")


# === FOOTER ===
st.markdown("""
<div class="footer">
    <strong>AskToDario</strong> — People Also Ask Explorer<br/>
    Creato da <a href="https://www.webinfermento.it/consulente-seo/" target="_blank">Dario Ciraci</a>
    | Senior SEO Consultant<br/>
    I risultati vengono cachati per 1 ora per velocizzare le ricerche ripetute
</div>
""", unsafe_allow_html=True)
